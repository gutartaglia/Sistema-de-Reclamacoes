import os
import io
from html.parser import HTMLParser

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches
from docx.image.image import Image as ImagemDocx

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CAMINHO_MODELO = os.path.join(BASE_DIR, "modelo", "Ficha_Reclamacao_Modelo_Padrao_1.docx")
PASTA_STATIC = os.path.join(BASE_DIR, "static")

MARGEM_SEGURANCA_POLEGADAS = 0.3
DPI_PADRAO = 96
RESERVA_CABECALHO_POLEGADAS = 1.3
ESPACAMENTO_ENTRE_FOTOS_POLEGADAS = 0.15
ALTURA_MINIMA_POLEGADAS = 0.8


class RelatoParser(HTMLParser):

    TAGS_NEGRITO = ("b", "strong")
    TAGS_ITALICO = ("i", "em")
    TAGS_QUEBRA = ("p", "div")

    def __init__(self):
        super().__init__()
        self.segmentos = []
        self.negrito = 0
        self.italico = 0
        self.marcado = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.TAGS_NEGRITO:
            self.negrito += 1
        elif tag in self.TAGS_ITALICO:
            self.italico += 1
        elif tag == "mark":
            self.marcado += 1
        elif tag == "br":
            self.segmentos.append(("\n", False, False, False))
        elif tag in self.TAGS_QUEBRA and self.segmentos:
            self.segmentos.append(("\n", False, False, False))

    def handle_endtag(self, tag):
        if tag in self.TAGS_NEGRITO:
            self.negrito = max(0, self.negrito - 1)
        elif tag in self.TAGS_ITALICO:
            self.italico = max(0, self.italico - 1)
        elif tag == "mark":
            self.marcado = max(0, self.marcado - 1)

    def handle_data(self, dados):
        if dados:
            self.segmentos.append((dados, self.negrito > 0, self.italico > 0, self.marcado > 0))


def _limpar_paragrafo(paragrafo):
    for run in list(paragrafo.runs):
        run.text = ""


def _preencher_celula(tabela, linha, coluna, linhas):
    linhas = [str(texto) for texto in linhas if texto]
    if not linhas:
        return

    celula = tabela.cell(linha, coluna)
    paragrafos_disponiveis = celula.paragraphs[1:]

    for indice, texto in enumerate(linhas):
        if indice < len(paragrafos_disponiveis):
            paragrafo = paragrafos_disponiveis[indice]
        else:
            paragrafo = celula.add_paragraph()

        _limpar_paragrafo(paragrafo)
        if paragrafo.runs:
            paragrafo.runs[0].text = texto
        else:
            paragrafo.add_run(texto)


def _preencher_relato(tabela, linha, coluna, relato_html):
    parser = RelatoParser()
    parser.feed(relato_html or "")
    segmentos = parser.segmentos

    if not any(texto.strip() for texto, *_ in segmentos if texto != "\n"):
        return

    celula = tabela.cell(linha, coluna)
    paragrafos = celula.paragraphs
    paragrafo = paragrafos[1] if len(paragrafos) > 1 else celula.add_paragraph()
    _limpar_paragrafo(paragrafo)

    for texto, negrito, italico, marcado in segmentos:
        if texto == "\n":
            paragrafo.add_run().add_break()
            continue

        run = paragrafo.add_run(texto)
        run.bold = negrito
        run.italic = italico
        if marcado:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _marcar_checkbox_fotos(documento, tem_fotos):
    alvo = "SIM" if tem_fotos else "NÃO"

    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    if "☐ SIM" in paragrafo.text and "☐ NÃO" in paragrafo.text:
                        novo_texto = paragrafo.text.replace(f"☐ {alvo}", f"☑ {alvo}")
                        _limpar_paragrafo(paragrafo)
                        if paragrafo.runs:
                            paragrafo.runs[0].text = novo_texto
                        else:
                            paragrafo.add_run(novo_texto)
                        return


def _todos_paragrafos(documento):
    for paragrafo in documento.paragraphs:
        yield paragrafo
    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    yield paragrafo


def _substituir_numero(documento, numero):
    for paragrafo in _todos_paragrafos(documento):
        if "____________" in paragrafo.text:
            for run in paragrafo.runs:
                if "____________" in run.text:
                    run.text = run.text.replace("____________", numero or "")


def _preencher_ocorrencia_referencia(tabela, numero, cliente, preferencia_cliente, nota_fiscal):
    partes = [f"Ficha Nº {numero}" if numero else None, cliente or None, preferencia_cliente or None]
    if nota_fiscal:
        partes.append(f"NF {nota_fiscal}")

    texto = " — ".join(parte for parte in partes if parte)
    _preencher_celula(tabela, 0, 0, [texto])


def _largura_ajustada(caminho_absoluto, largura_maxima_polegadas, altura_maxima_polegadas):
    try:
        imagem = ImagemDocx.from_file(caminho_absoluto)
        dpi_horizontal = imagem.horz_dpi or DPI_PADRAO
        dpi_vertical = imagem.vert_dpi or DPI_PADRAO
        largura_nativa_polegadas = imagem.px_width / dpi_horizontal
        altura_nativa_polegadas = imagem.px_height / dpi_vertical
    except Exception:
        return Inches(largura_maxima_polegadas)

    fator = min(
        1.0,
        largura_maxima_polegadas / largura_nativa_polegadas,
        altura_maxima_polegadas / altura_nativa_polegadas
    )

    return Inches(largura_nativa_polegadas * fator)


def _altura_maxima_por_foto(secao, quantidade_fotos):
    altura_pagina_util = secao.page_height.inches - secao.top_margin.inches - secao.bottom_margin.inches
    orcamento_disponivel = max(2.0, altura_pagina_util - RESERVA_CABECALHO_POLEGADAS)

    if quantidade_fotos <= 1:
        return orcamento_disponivel / 2

    espacamento_total = ESPACAMENTO_ENTRE_FOTOS_POLEGADAS * (quantidade_fotos - 1)
    return max(ALTURA_MINIMA_POLEGADAS, (orcamento_disponivel - espacamento_total) / quantidade_fotos)


def _inserir_fotos_investigacao(tabela, fotos, secao):
    fotos_validas = [
        (id_foto, caminho) for id_foto, caminho in (fotos or [])
        if os.path.exists(os.path.join(PASTA_STATIC, caminho))
    ]

    if not fotos_validas:
        return

    celula = tabela.cell(0, 0)

    for paragrafo_extra in celula.paragraphs[1:]:
        elemento = paragrafo_extra._element
        elemento.getparent().remove(elemento)

    largura_disponivel = celula.width.inches if celula.width else 6.5
    largura_maxima = max(1.0, largura_disponivel - MARGEM_SEGURANCA_POLEGADAS)
    altura_maxima = _altura_maxima_por_foto(secao, len(fotos_validas))

    primeira_foto = True

    for _, caminho in fotos_validas:
        caminho_absoluto = os.path.join(PASTA_STATIC, caminho)

        if primeira_foto:
            paragrafo = celula.paragraphs[0]
            _limpar_paragrafo(paragrafo)
            primeira_foto = False
        else:
            paragrafo = celula.add_paragraph()

        try:
            largura = _largura_ajustada(caminho_absoluto, largura_maxima, altura_maxima)
            paragrafo.add_run().add_picture(caminho_absoluto, width=largura)
        except Exception:
            continue


def _preencher_acoes_corretivas(tabela, acoes):
    if not acoes:
        return

    celula = tabela.cell(0, 0)

    for paragrafo_extra in celula.paragraphs[1:]:
        elemento = paragrafo_extra._element
        elemento.getparent().remove(elemento)

    primeiro_paragrafo = True
    total = len(acoes)

    for indice, (_, titulo, descricao, _concluida) in enumerate(acoes, start=1):

        if primeiro_paragrafo:
            paragrafo_numero = celula.paragraphs[0]
            _limpar_paragrafo(paragrafo_numero)
            primeiro_paragrafo = False
        else:
            paragrafo_numero = celula.add_paragraph()

        run_numero = paragrafo_numero.add_run(f"Ação #{indice}")
        run_numero.bold = True

        paragrafo_titulo = celula.add_paragraph()
        run_titulo = paragrafo_titulo.add_run(titulo or "")
        run_titulo.bold = True

        celula.add_paragraph(descricao or "")

        if indice < total:
            celula.add_paragraph()


def _formatar_quantidade_produto(quantidade):
    try:
        quantidade = float(quantidade)
    except (TypeError, ValueError):
        return str(quantidade or "")

    if quantidade.is_integer():
        return str(int(quantidade))
    return str(quantidade)


def _preencher_produtos_defeituosos(tabela, produtos):
    linhas = []
    for _, _origem, codigo, descricao, quantidade in (produtos or []):
        rotulo = " — ".join(parte for parte in (codigo, descricao) if parte)
        linhas.append(f"{_formatar_quantidade_produto(quantidade)}x {rotulo}".strip())

    _preencher_celula(tabela, 1, 1, linhas)


def gerar_ficha_reclamacao(rdc, lista_fotos=None, acoes=None, produtos_defeituosos=None):
    (
        id_, numero, cliente, email, cpf_cnpj, telefone, endereco, cep,
        nota_fiscal, data, vendedor, fotos, relato, quantidade_defeitos,
        participantes, preferencia_cliente, causa_principal, criado_em, unidade
    ) = rdc

    documento = docx.Document(CAMINHO_MODELO)

    _substituir_numero(documento, numero)

    tabela_dados = documento.tables[1]

    linha_cliente = f"{cliente} — CPF/CNPJ: {cpf_cnpj}" if cpf_cnpj else cliente
    _preencher_celula(tabela_dados, 0, 0, [linha_cliente])

    contato = []
    if email:
        for endereco_email in email.split(";"):
            endereco_email = endereco_email.strip()
            if endereco_email:
                contato.append(endereco_email)
    if telefone:
        contato.append(f"Telefone: {telefone}")
    _preencher_celula(tabela_dados, 0, 1, contato)

    endereco_linhas = [endereco] if endereco else []
    if cep:
        endereco_linhas.append(f"CEP: {cep}")
    _preencher_celula(tabela_dados, 1, 0, endereco_linhas)

    if produtos_defeituosos:
        _preencher_produtos_defeituosos(tabela_dados, produtos_defeituosos)
    else:
        _preencher_celula(tabela_dados, 1, 1, [quantidade_defeitos] if quantidade_defeitos else [])

    nf_linha = f"NF: {nota_fiscal or '-'} | Data: {data or '-'} | Vendedor: {vendedor or '-'}"
    _preencher_celula(tabela_dados, 2, 0, [nf_linha])

    tabela_extra = documento.tables[2]
    _preencher_celula(tabela_extra, 0, 0, [participantes] if participantes else [])
    _preencher_celula(tabela_extra, 1, 0, [preferencia_cliente] if preferencia_cliente else [])
    _preencher_relato(tabela_extra, 2, 0, relato)

    _marcar_checkbox_fotos(documento, bool(fotos))

    tabela_ocorrencia_investigacao = documento.tables[3]
    tabela_ocorrencia_acoes = documento.tables[5]
    tabela_investigacao_conteudo = documento.tables[4]
    tabela_acoes_conteudo = documento.tables[6]

    _preencher_ocorrencia_referencia(
        tabela_ocorrencia_investigacao, numero, cliente, preferencia_cliente, nota_fiscal
    )
    _preencher_ocorrencia_referencia(
        tabela_ocorrencia_acoes, numero, cliente, preferencia_cliente, nota_fiscal
    )

    _inserir_fotos_investigacao(tabela_investigacao_conteudo, lista_fotos, documento.sections[0])
    _preencher_acoes_corretivas(tabela_acoes_conteudo, acoes)

    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
