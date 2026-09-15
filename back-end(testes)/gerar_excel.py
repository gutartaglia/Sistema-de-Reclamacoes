import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.image import Image as ImagemExcel
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CAMINHO_LOGO = os.path.join(BASE_DIR, "static", "img", "logo_santa_fe.png")
ALTURA_LOGO_PX = 72

FONTE = "Arial"

COR_TITULO = "1F3864"
COR_SUBTITULO = "595959"
COR_SECAO_FUNDO = "1F3864"
COR_SECAO_TEXTO = "FFFFFF"
COR_CABECALHO_TABELA = "41729F"
COR_CABECALHO_TEXTO = "FFFFFF"
COR_LINHA_DADOS_FUNDO = "F2F2F2"
COR_LINHA_DADOS_TEXTO = "222222"
COR_TOTAL_FUNDO = "1F3864"
COR_TOTAL_TEXTO = "FFFFFF"
COR_BORDA = "D9D9D9"

# Escala monocromática 1 do Excel (tons de azul, derivados do Accent 1 do tema Office):
# do mais claro ao mais escuro — DAE3F3, B4C7E7, 8EA9DB, 4472C4, 2E5395, 1F3864.
# Exceções mantidas fora da escala, para destaque/legibilidade: o amarelo/dourado da
# linha de % acumulado do Pareto e a paleta colorida do gráfico de causas (várias
# causas lado a lado ficam difíceis de distinguir só com tons de azul).
COR_LINHA_MENSAL = "2E5395"
COR_INDICE = "4472C4"
COR_RDC_ANO = "1F3864"
COR_PARETO_BARRA = "4472C4"
COR_PARETO_LINHA = "FFC000"
COR_FILIAL = "8EA9DB"
COR_MATRIZ = "1F3864"
COR_TOTAL_LINHA = "4472C4"

PALETA_CAUSAS = [
    "2563EB", "16A34A", "DC2626", "F59E0B", "8B5CF6", "06B6D4",
    "EC4899", "84CC16", "F97316", "14B8A6", "A855F7", "EAB308",
    "EF4444", "22C55E", "3B82F6", "D946EF", "0EA5E9", "FACC15",
    "F43F5E", "10B981", "6366F1", "FB923C", "C026D3", "65A30D",
]

BORDA_FINA = Border(
    top=Side(style="thin", color=COR_BORDA),
    bottom=Side(style="thin", color=COR_BORDA),
)

ALINHAR_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALINHAR_ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)

ALTURA_LINHA_CM = 0.53


def _linhas_reservadas(altura_cm, margem=3):
    """Calcula quantas linhas reservar após um gráfico, de acordo com a altura real dele."""
    return int(altura_cm / ALTURA_LINHA_CM) + margem


def _posicionar_eixos(eixo_x, eixo_y, posicao_x="b", posicao_y="l"):
    """O openpyxl cria os eixos com axPos='l' por padrão, mesmo em gráficos de coluna —
    isso faz o Excel não desenhar as categorias (meses/anos) embaixo das barras."""
    eixo_x.axPos = posicao_x
    eixo_x.delete = False
    eixo_y.axPos = posicao_y
    eixo_y.delete = False


def _remover_numeros_eixo(eixo_y):
    """Os valores já aparecem nos rótulos em cima das barras/linha, então os números
    do eixo lateral ficam redundantes. Usamos um formato de número em branco (;;;)
    em vez de eixo.delete=True: apagar o eixo também apaga o espaço que ele reserva,
    e o título do eixo e a legenda embaixo colam nas linhas de grade. Com o formato em
    branco o eixo continua "existindo" (mantém o respiro do layout), só o texto some.
    Sem as aspas ao redor: com aspas o Excel exibe ";;;" literalmente; os 4 segmentos
    vazios (positivo;negativo;zero;texto) é que fazem o valor não aparecer."""
    eixo_y.numFmt = ';;;'
    eixo_y.delete = False
    eixo_y.majorGridlines = None
    eixo_y.minorGridlines = None


def _sem_contorno(grafico):
    """Remove o contorno (borda) da área do gráfico e da área de plotagem, e garante
    que nenhum eixo desenhe linhas de grade — evitando que rótulos de dados fiquem
    sobrepostos a uma linha de grade."""
    sem_linha = GraphicalProperties(ln=LineProperties(noFill=True))
    grafico.graphical_properties = sem_linha
    grafico.plot_area.graphicalProperties = sem_linha
    for eixo in (grafico.x_axis, grafico.y_axis):
        eixo.majorGridlines = None
        eixo.minorGridlines = None


def _folga_eixo_y(eixo_y, maior_valor, folga=0.18):
    """Dá uma margem acima da maior barra/ponto para o rótulo de valor não encostar no título."""
    if not maior_valor:
        return
    eixo_y.scaling.max = maior_valor * (1 + folga)


def _rotulos_dados(posicao=None, formato_numero=None, tamanho_fonte=8, cor_fonte="404040"):
    rotulos = DataLabelList()
    rotulos.showVal = True
    rotulos.showCatName = False
    rotulos.showSerName = False
    rotulos.showLegendKey = False
    rotulos.showPercent = False
    rotulos.showBubbleSize = False
    if posicao:
        rotulos.dLblPos = posicao
    if formato_numero:
        rotulos.numFmt = formato_numero
    rotulos.txPr = _texto_rotulo(tamanho_fonte, cor_fonte)
    return rotulos


def _texto_rotulo(tamanho_fonte, cor_fonte):
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import (
        Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties, Font as DrawingFont
    )

    propriedades = CharacterProperties(
        sz=tamanho_fonte * 100, b=True,
        solidFill=cor_fonte, latin=DrawingFont(typeface=FONTE)
    )
    paragrafo = Paragraph(pPr=ParagraphProperties(defRPr=propriedades), endParaRPr=propriedades)
    return RichText(bodyPr=RichTextProperties(), p=[paragrafo])


def _cor_causa(indice):
    return PALETA_CAUSAS[indice % len(PALETA_CAUSAS)]


def _fonte(tamanho, negrito=False, cor="000000"):
    return Font(name=FONTE, size=tamanho, bold=negrito, color=cor)


def _preencher(cor):
    return PatternFill("solid", fgColor=cor)


def _inserir_logo(ws):
    if not os.path.exists(CAMINHO_LOGO):
        return
    logo = ImagemExcel(CAMINHO_LOGO)
    proporcao = logo.width / logo.height
    logo.height = ALTURA_LOGO_PX
    logo.width = ALTURA_LOGO_PX * proporcao
    ws.add_image(logo, "A1")


def _titulo(ws, titulo, subtitulo, num_colunas):
    ultima_coluna = get_column_letter(max(num_colunas + 1, 7))
    ws.merge_cells(f"B1:{ultima_coluna}2")
    ws.merge_cells(f"B3:{ultima_coluna}3")

    _inserir_logo(ws)

    celula_titulo = ws["B1"]
    celula_titulo.value = titulo
    celula_titulo.font = _fonte(20, negrito=True, cor=COR_TITULO)
    celula_titulo.alignment = Alignment(horizontal="right", vertical="center")

    celula_subtitulo = ws["B3"]
    celula_subtitulo.value = subtitulo
    celula_subtitulo.font = _fonte(11, cor=COR_SUBTITULO)
    celula_subtitulo.alignment = Alignment(horizontal="right", vertical="center")

    ws.row_dimensions[1].height = 19.5
    ws.row_dimensions[2].height = 25.5
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 9.75

    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 28
    for indice in range(3, num_colunas + 2):
        ws.column_dimensions[get_column_letter(indice)].width = 17


def _rodape(ws, linha, nome_relatorio):
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells(f"B{linha}:F{linha}")
    celula = ws.cell(row=linha, column=2)
    celula.value = f"Documento gerado automaticamente em {agora} — Setor de Qualidade Santa Fé — Relatório: {nome_relatorio}."
    celula.font = _fonte(9, cor=COR_SUBTITULO)
    celula.alignment = ALINHAR_ESQUERDA


def _escrever_secao_tabela(ws, linha_inicio, numero, titulo, colunas, linhas, resumo_formulas, resumo_rotulo="Total"):
    num_colunas = len(colunas)
    ultima_coluna_letra = get_column_letter(num_colunas + 1)

    ws.merge_cells(f"B{linha_inicio}:{ultima_coluna_letra}{linha_inicio}")
    celula_secao = ws.cell(row=linha_inicio, column=2, value=f"{numero}. {titulo}")
    celula_secao.font = _fonte(13, negrito=True, cor=COR_SECAO_TEXTO)
    celula_secao.fill = _preencher(COR_SECAO_FUNDO)
    celula_secao.alignment = ALINHAR_ESQUERDA
    for indice in range(2, num_colunas + 2):
        ws.cell(row=linha_inicio, column=indice).fill = _preencher(COR_SECAO_FUNDO)

    linha_cabecalho = linha_inicio + 1
    for indice, texto in enumerate(colunas, start=2):
        celula = ws.cell(row=linha_cabecalho, column=indice, value=texto)
        celula.font = _fonte(10, negrito=True, cor=COR_CABECALHO_TEXTO)
        celula.fill = _preencher(COR_CABECALHO_TABELA)
        celula.alignment = ALINHAR_CENTRO
        celula.border = BORDA_FINA

    linha_atual = linha_cabecalho + 1
    for linha_dados in linhas:
        for indice, valor in enumerate(linha_dados, start=2):
            celula = ws.cell(row=linha_atual, column=indice, value=valor)
            celula.font = _fonte(10, cor=COR_LINHA_DADOS_TEXTO)
            celula.fill = _preencher(COR_LINHA_DADOS_FUNDO)
            celula.alignment = ALINHAR_CENTRO if indice > 2 else ALINHAR_ESQUERDA
            celula.border = BORDA_FINA
        linha_atual += 1

    linha_resumo = linha_atual
    celula_rotulo = ws.cell(row=linha_resumo, column=2, value=resumo_rotulo)
    celula_rotulo.font = _fonte(10, negrito=True, cor=COR_TOTAL_TEXTO)
    celula_rotulo.fill = _preencher(COR_TOTAL_FUNDO)
    celula_rotulo.alignment = ALINHAR_ESQUERDA
    celula_rotulo.border = BORDA_FINA

    for indice, formula in enumerate(resumo_formulas, start=3):
        celula = ws.cell(row=linha_resumo, column=indice, value=formula)
        celula.font = _fonte(10, negrito=True, cor=COR_TOTAL_TEXTO)
        celula.fill = _preencher(COR_TOTAL_FUNDO)
        celula.alignment = ALINHAR_CENTRO
        celula.border = BORDA_FINA

    return linha_cabecalho, linha_atual - 1, linha_resumo


def _formula_soma(ws_nome, coluna_letra, primeira_linha, ultima_linha):
    return f"=SUM({coluna_letra}{primeira_linha}:{coluna_letra}{ultima_linha})"


def _formula_media(ws_nome, coluna_letra, primeira_linha, ultima_linha):
    return f"=AVERAGE({coluna_letra}{primeira_linha}:{coluna_letra}{ultima_linha})"


def _secao_mensal(wb, ws, linha_inicio, numero, mensal, cor_linha):
    if not mensal:
        return linha_inicio

    colunas = ["Categoria", "Reclamações"]
    linhas = [[item["mes_label"], item["total"]] for item in mensal]
    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, "Evolução Mensal de Reclamações", colunas, linhas,
        [_formula_soma(ws.title, "C", linha_inicio + 2, linha_inicio + 1 + len(linhas))]
    )

    dados_ref = Reference(ws, min_col=3, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    grafico = LineChart()
    grafico.title = "Reclamações por Mês"
    grafico.add_data(dados_ref, titles_from_data=True)
    grafico.set_categories(cat_ref)
    grafico.series[0].graphicalProperties.line.solidFill = cor_linha
    grafico.series[0].graphicalProperties.line.width = 25000
    grafico.dataLabels = _rotulos_dados("t")
    grafico.legend = None
    _posicionar_eixos(grafico.x_axis, grafico.y_axis)
    _folga_eixo_y(grafico.y_axis, max(item["total"] for item in mensal))
    _remover_numeros_eixo(grafico.y_axis)
    _sem_contorno(grafico)
    grafico.height = 7.5
    grafico.width = 16
    ws.add_chart(grafico, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(grafico.height)


def _secao_indice(wb, ws, linha_inicio, numero, mensal, cor_barra):
    if not mensal:
        return linha_inicio

    colunas = ["Categoria", "Índice (nº / 2000)"]
    linhas = [[item["mes_label"], item["indice"]] for item in mensal]
    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, "Índice de Reclamações (nº / 2000 unidades)", colunas, linhas,
        [_formula_media(ws.title, "C", linha_inicio + 2, linha_inicio + 1 + len(linhas))],
        resumo_rotulo="Média"
    )

    for linha in range(linha_cab + 1, linha_fim + 1):
        ws.cell(row=linha, column=3).number_format = "0.0000"
    ws.cell(row=linha_resumo, column=3).number_format = "0.0000"

    dados_ref = Reference(ws, min_col=3, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    grafico = BarChart()
    grafico.type = "col"
    grafico.title = "Índice de Reclamações por Mês"
    grafico.add_data(dados_ref, titles_from_data=True)
    grafico.set_categories(cat_ref)
    grafico.series[0].graphicalProperties.solidFill = cor_barra
    grafico.dataLabels = _rotulos_dados("outEnd", formato_numero="0.0000")
    grafico.legend = None
    _posicionar_eixos(grafico.x_axis, grafico.y_axis)
    _folga_eixo_y(grafico.y_axis, max(item["indice"] for item in mensal))
    _remover_numeros_eixo(grafico.y_axis)
    _sem_contorno(grafico)
    grafico.height = 7.5
    grafico.width = 16
    ws.add_chart(grafico, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(grafico.height)


def _secao_causas_agrupadas(wb, ws, linha_inicio, numero, titulo, rotulo_categoria, causas_lista, tabela, campo_categoria):
    if not tabela or not causas_lista:
        return linha_inicio

    colunas = [rotulo_categoria] + causas_lista + ["Total"]
    linhas = []
    for item in tabela:
        linha = [item[campo_categoria]] + list(item["valores"])
        linhas.append(linha)

    ultima_col_causa_letra = get_column_letter(len(causas_lista) + 2)
    total_col_letra = get_column_letter(len(causas_lista) + 3)

    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, titulo, colunas, linhas,
        [_formula_soma(ws.title, get_column_letter(c), linha_inicio + 2, linha_inicio + 1 + len(linhas))
         for c in range(3, len(causas_lista) + 3)]
    )

    for indice_linha in range(linha_cab + 1, linha_fim + 1):
        celula_total = ws.cell(row=indice_linha, column=len(causas_lista) + 3)
        celula_total.value = f"=SUM(C{indice_linha}:{ultima_col_causa_letra}{indice_linha})"
        celula_total.font = _fonte(10, negrito=True, cor=COR_LINHA_DADOS_TEXTO)
        celula_total.fill = _preencher(COR_LINHA_DADOS_FUNDO)
        celula_total.alignment = ALINHAR_CENTRO
        celula_total.border = BORDA_FINA

    celula_resumo_total = ws.cell(row=linha_resumo, column=len(causas_lista) + 3)
    celula_resumo_total.value = f"=SUM({total_col_letra}{linha_cab + 1}:{total_col_letra}{linha_fim})"
    celula_resumo_total.font = _fonte(10, negrito=True, cor=COR_TOTAL_TEXTO)
    celula_resumo_total.fill = _preencher(COR_TOTAL_FUNDO)
    celula_resumo_total.alignment = ALINHAR_CENTRO
    celula_resumo_total.border = BORDA_FINA

    dados_ref = Reference(ws, min_col=3, max_col=len(causas_lista) + 2, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    numero_causas = len(causas_lista)

    grafico = BarChart()
    grafico.type = "col"
    grafico.grouping = "clustered"
    grafico.title = titulo
    grafico.add_data(dados_ref, titles_from_data=True)
    grafico.set_categories(cat_ref)

    for indice, serie in enumerate(grafico.series):
        serie.graphicalProperties.solidFill = _cor_causa(indice)

    grafico.dataLabels = _rotulos_dados("outEnd")

    _posicionar_eixos(grafico.x_axis, grafico.y_axis)
    _folga_eixo_y(grafico.y_axis, max(valor for linha in linhas for valor in linha[1:]))
    _remover_numeros_eixo(grafico.y_axis)
    _sem_contorno(grafico)

    # Tamanho: a largura acompanha as categorias do eixo X (meses/anos), não a
    # quantidade de causas — isso evita gráficos gigantes quando há muitas causas.
    grafico.width = min(28, max(18, 14 + len(linhas) * 0.6))

    # Legenda embaixo, com altura extra proporcional ao nº de linhas que ela vai ocupar,
    # para nunca sobrepor as barras.
    entradas_por_linha_legenda = 6
    linhas_legenda = -(-numero_causas // entradas_por_linha_legenda)
    grafico.height = min(18, max(9, 8 + linhas_legenda * 1.1))

    grafico.legend.position = "b"
    grafico.legend.overlay = False

    ws.add_chart(grafico, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(grafico.height)


def _secao_pareto(wb, ws, linha_inicio, numero, titulo, pareto):
    if not pareto:
        return linha_inicio

    colunas = ["Causa", "Frequência", "Freq. acumulada", "% do total", "% acumulado"]
    linhas = [
        [item["causa"], item["frequencia"], item["frequencia_acumulada"], item["percentual"], item["percentual_acumulado"]]
        for item in pareto
    ]

    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, titulo, colunas, linhas,
        [_formula_soma(ws.title, "C", linha_inicio + 2, linha_inicio + 1 + len(linhas))]
    )

    for linha in range(linha_cab + 1, linha_fim + 1):
        ws.cell(row=linha, column=5).number_format = "0.00"
        ws.cell(row=linha, column=6).number_format = "0.00"

    dados_ref = Reference(ws, min_col=3, min_row=linha_cab, max_row=linha_fim)
    pct_ref = Reference(ws, min_col=6, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    barras = BarChart()
    barras.type = "col"
    barras.title = titulo
    barras.add_data(dados_ref, titles_from_data=True)
    barras.set_categories(cat_ref)
    barras.series[0].graphicalProperties.solidFill = COR_PARETO_BARRA
    barras.series[0].dLbls = _rotulos_dados("outEnd")
    barras.y_axis.title = "Frequência"
    barras.y_axis.crosses = "min"
    _posicionar_eixos(barras.x_axis, barras.y_axis)
    _folga_eixo_y(barras.y_axis, max(item["frequencia"] for item in pareto))
    _remover_numeros_eixo(barras.y_axis)

    linha_grafico = LineChart()
    linha_grafico.add_data(pct_ref, titles_from_data=True)
    linha_grafico.series[0].graphicalProperties.line.solidFill = COR_PARETO_LINHA
    linha_grafico.series[0].graphicalProperties.line.width = 25000
    linha_grafico.series[0].dLbls = _rotulos_dados("t", formato_numero='0.0"%"', cor_fonte=COR_PARETO_LINHA)
    linha_grafico.y_axis.axId = 200
    linha_grafico.y_axis.title = "% acumulado"
    linha_grafico.y_axis.crosses = "max"
    linha_grafico.y_axis.axPos = "r"
    linha_grafico.y_axis.delete = False
    _folga_eixo_y(linha_grafico.y_axis, 100)
    _remover_numeros_eixo(linha_grafico.y_axis)

    barras += linha_grafico
    barras.legend.position = "b"
    barras.legend.overlay = False
    _sem_contorno(barras)
    barras.height = 10
    barras.width = min(26, max(18, 14 + len(linhas) * 0.6))
    ws.add_chart(barras, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(barras.height)


def _secao_rdc_por_ano(wb, ws, linha_inicio, numero, titulo, rdc_por_ano, cor_barra):
    if not rdc_por_ano:
        return linha_inicio

    colunas = ["Categoria", "Total de RDCs"]
    linhas = [[str(ano), total] for ano, total in rdc_por_ano]

    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, titulo, colunas, linhas,
        [_formula_soma(ws.title, "C", linha_inicio + 2, linha_inicio + 1 + len(linhas))]
    )

    dados_ref = Reference(ws, min_col=3, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    grafico = BarChart()
    grafico.type = "col"
    grafico.title = titulo
    grafico.add_data(dados_ref, titles_from_data=True)
    grafico.set_categories(cat_ref)
    grafico.series[0].graphicalProperties.solidFill = cor_barra
    grafico.dataLabels = _rotulos_dados("outEnd")
    grafico.legend = None
    _posicionar_eixos(grafico.x_axis, grafico.y_axis)
    _folga_eixo_y(grafico.y_axis, max(total for _, total in rdc_por_ano))
    _remover_numeros_eixo(grafico.y_axis)
    _sem_contorno(grafico)
    grafico.height = 7.5
    grafico.width = 16
    ws.add_chart(grafico, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(grafico.height)


def _secao_mensal_combinado(wb, ws, linha_inicio, numero, mensal):
    if not mensal:
        return linha_inicio

    colunas = ["Mês", "Filial", "Matriz", "Total"]
    linhas = [[item["mes_label"], item["filial"], item["matriz"], item["total"]] for item in mensal]

    linha_cab, linha_fim, linha_resumo = _escrever_secao_tabela(
        ws, linha_inicio, numero, "Reclamações por Mês — Filial x Matriz", colunas, linhas,
        [
            _formula_soma(ws.title, "C", linha_inicio + 2, linha_inicio + 1 + len(linhas)),
            _formula_soma(ws.title, "D", linha_inicio + 2, linha_inicio + 1 + len(linhas)),
            _formula_soma(ws.title, "E", linha_inicio + 2, linha_inicio + 1 + len(linhas)),
        ]
    )

    barras_ref = Reference(ws, min_col=3, max_col=4, min_row=linha_cab, max_row=linha_fim)
    total_ref = Reference(ws, min_col=5, min_row=linha_cab, max_row=linha_fim)
    cat_ref = Reference(ws, min_col=2, min_row=linha_cab + 1, max_row=linha_fim)

    barras = BarChart()
    barras.type = "col"
    barras.grouping = "clustered"
    barras.title = "Reclamações por Mês — Filial x Matriz"
    barras.add_data(barras_ref, titles_from_data=True)
    barras.set_categories(cat_ref)
    barras.series[0].graphicalProperties.solidFill = COR_FILIAL
    barras.series[0].dLbls = _rotulos_dados("outEnd")
    barras.series[1].graphicalProperties.solidFill = COR_MATRIZ
    barras.series[1].dLbls = _rotulos_dados("outEnd")

    linha_grafico = LineChart()
    linha_grafico.add_data(total_ref, titles_from_data=True)
    linha_grafico.series[0].graphicalProperties.line.solidFill = COR_TOTAL_LINHA
    linha_grafico.series[0].graphicalProperties.line.width = 25000
    linha_grafico.series[0].dLbls = _rotulos_dados("t", cor_fonte=COR_TOTAL_LINHA)

    _posicionar_eixos(barras.x_axis, barras.y_axis)
    maior_valor = max(max(item["filial"], item["matriz"], item["total"]) for item in mensal)
    _folga_eixo_y(barras.y_axis, maior_valor)
    _remover_numeros_eixo(barras.y_axis)

    barras += linha_grafico
    barras.legend.position = "b"
    barras.legend.overlay = False
    _sem_contorno(barras)
    barras.height = 10
    barras.width = min(26, max(18, 14 + len(linhas) * 0.6))
    ws.add_chart(barras, f"B{linha_resumo + 2}")

    return linha_resumo + 2 + _linhas_reservadas(barras.height)


def gerar_excel_unidade(nome_unidade, mensal, causas_lista, causas_tabela, pareto, rdc_por_ano, ano_filtro="todos"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    num_colunas_titulo = max(len(causas_lista) + 2, 5)
    _titulo(ws, "Dashboard de Reclamações de Clientes", f"Indicadores de Qualidade — {nome_unidade}", num_colunas_titulo)

    linha = 6
    numero = 1

    if mensal:
        linha = _secao_mensal(wb, ws, linha, numero, mensal, COR_LINHA_MENSAL)
        numero += 1

    if mensal:
        linha = _secao_indice(wb, ws, linha, numero, mensal, COR_INDICE)
        numero += 1

    if causas_tabela and causas_lista:
        linha = _secao_causas_agrupadas(
            wb, ws, linha, numero, "Causas das Reclamações por Mês", "Mês", causas_lista, causas_tabela, "mes_label"
        )
        numero += 1

    if pareto:
        linha = _secao_pareto(wb, ws, linha, numero, "Diagrama de Pareto das Causas", pareto)
        numero += 1

    if rdc_por_ano:
        linha = _secao_rdc_por_ano(wb, ws, linha, numero, "Total de RDCs por Ano", rdc_por_ano, COR_RDC_ANO)
        numero += 1

    if numero == 1:
        ws.cell(row=6, column=2, value="Nenhum dado disponível para este filtro.")

    filtro_texto = "todos os anos" if ano_filtro == "todos" else f"ano {ano_filtro}"
    _rodape(ws, linha + 1, f"{nome_unidade} ({filtro_texto})")

    return wb


def gerar_excel_combinado(causas_top, causas_ano_tabela, pareto, rdc_por_ano, mensal, ano_filtro="todos"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    num_colunas_titulo = max(len(causas_top) + 2, 5)
    _titulo(ws, "Dashboard de Reclamações de Clientes", "Consolidado de Indicadores de Qualidade — Filial + Matriz", num_colunas_titulo)

    linha = 6
    numero = 1

    if mensal:
        linha = _secao_mensal_combinado(wb, ws, linha, numero, mensal)
        numero += 1

    if causas_top and causas_ano_tabela:
        linha = _secao_causas_agrupadas(
            wb, ws, linha, numero, "4 Causas Principais por Ano", "Ano", causas_top, causas_ano_tabela, "ano"
        )
        numero += 1

    if pareto:
        linha = _secao_pareto(wb, ws, linha, numero, "Diagrama de Pareto das Causas — Filial + Matriz", pareto)
        numero += 1

    if rdc_por_ano:
        linha = _secao_rdc_por_ano(
            wb, ws, linha, numero, "Total de RDCs por Ano — Filial + Matriz", rdc_por_ano, COR_RDC_ANO
        )
        numero += 1

    if numero == 1:
        ws.cell(row=6, column=2, value="Nenhum dado disponível para este filtro.")

    filtro_texto = "todos os anos" if ano_filtro == "todos" else f"ano {ano_filtro}"
    _rodape(ws, linha + 1, f"Filial + Matriz ({filtro_texto})")

    return wb
